# 文件作用：测试报告读取、过滤、JSON 与 Excel 导出和标准答案匹配。
# 为什么有它：避免展示丢失信息、误把文本当公式或读取不应访问的文件。
"""Read-only export contracts; network access is forbidden."""

import hashlib
import json
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from src.comparison.demo import simulated_report
from src.evaluation.aggregation import aggregate_results
from src.report import ReportError, excel_bytes, export_report, json_bytes, load_report, parse_report
from src.report.data import expected_answers, filter_cases, mode_notice, report_catalog


# 做什么：阻止同步与异步 HTTP 调用。
# 为什么需要：读取和导出报告不应访问模型或其他网络服务。
@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    # 做什么：出现网络发送就让测试失败。
    # 为什么需要：保证报告功能保持只读本地数据。
    def forbidden(*args, **kwargs):
        raise AssertionError("Report export must never call LLM or network")
    monkeypatch.setattr(httpx.Client, "send", forbidden)
    monkeypatch.setattr(httpx.AsyncClient, "send", forbidden)


# 做什么：构造覆盖正式题目结构的模拟报告。
# 为什么需要：给导出测试提供完整但不依赖真实 API 的输入。
@pytest.fixture
def report():
    return simulated_report(Path("data/corpus/ecommerce_v1.json"), Path("data/datasets/ecommerce_eval_v1.json"), 3)


# 做什么：验证完整报告 JSON 往返保留逐题和原始评分信息。
# 为什么需要：导出应是完整证据而非只保存汇总。
def test_full_json_roundtrip_preserves_nested_results(report):
    assert parse_report(json_bytes(report)) == report
    payload = json.loads(json_bytes(report))
    assert len(payload["case_results"]) == 60
    assert payload["case_results"][0]["judge_result"]["metadata"]["raw_output"]
    assert payload["judge_mode"] == "MOCK"


# 做什么：验证坏报告给出明确解析错误。
# 为什么需要：页面与导出不能接受不完整结构。
@pytest.mark.parametrize("content", [b"not JSON", b"{}", b"[]"])
def test_invalid_report_is_clear(content):
    with pytest.raises(ReportError, match="Invalid EvaluationReport"):
        parse_report(content)


# 做什么：验证生成 JSON、Excel 和四张规定工作表。
# 为什么需要：保证导出功能符合交付约定。
def test_export_generates_files_and_four_sheets(report, tmp_path):
    paths = export_report(report, tmp_path)
    assert paths[0].name == "evaluation_report.json"
    assert paths[1].name == "evaluation_report.xlsx"
    assert load_report(paths[0]) == report
    book = load_workbook(paths[1])
    assert book.sheetnames == ["Summary", "Case Results", "Retrieval Metrics", "Judge Results"]
    for sheet in book:
        assert "JUDGE MODE: MOCK" in sheet["A2"].value
        assert sheet.freeze_panes == "A5" and sheet.auto_filter.ref
    assert book["Case Results"].max_row == 64
    assert book["Case Results"]["F5"].value is not None
    assert book["Retrieval Metrics"]["D5"].value == report.case_results[0].retrieval_metrics.recall_at_k
    assert book["Judge Results"]["D5"].data_type == "n"
    assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
    summary = {row[1]: row[2] for row in book["Summary"].iter_rows(min_row=5, values_only=True)}
    assert summary["total_cases"] == 60
    assert summary["corpus_version"] == "sha256:" + report.corpus_sha256


# 做什么：验证公式样式文本、控制字符和长文本安全写入。
# 为什么需要：防止 Excel 执行字符串或生成损坏文件。
def test_excel_literal_formulas_controls_and_long_text(report):
    report.case_results[0].query = '=HYPERLINK("https://example.invalid", "click")'
    report.case_results[0].rag_result.rag_answer = "bad\x00control" + "答" * 40000
    book = load_workbook(BytesIO(excel_bytes(report)))
    assert book["Case Results"]["E5"].value.startswith("=HYPERLINK")
    assert book["Case Results"]["E5"].data_type == "s"
    value = book["Case Results"]["G5"].value
    assert "\\u0000" in value and "TRUNCATED" in value and len(value) <= 32767
    assert parse_report(json_bytes(report)).case_results[0].rag_result.rag_answer.endswith("答" * 40000)


# 做什么：验证缺裁判和被排除的检索分数保持空白。
# 为什么需要：不能把没有测量的内容显示成真实零分。
def test_missing_judge_and_excluded_retrieval_stay_blank(report):
    report.case_results[0].judge_result = None
    book = load_workbook(BytesIO(excel_bytes(report)))
    assert book["Judge Results"]["B5"].value == "UNAVAILABLE"
    assert book["Judge Results"]["D5"].value is None
    n = next(i for i, row in enumerate(report.case_results, 5) if not row.answerable)
    assert book["Retrieval Metrics"].cell(n, 4).value is None
    assert book["Retrieval Metrics"].cell(n, 9).value == "EXCLUDED"


# 做什么：验证标准答案只取自哈希一致的数据集。
# 为什么需要：避免新试卷答案被配到旧报告上。
def test_matching_answers_require_exact_hash(report, tmp_path):
    dataset = Path("data/datasets/ecommerce_eval_v1.json").read_bytes()
    (tmp_path / "matching.json").write_bytes(dataset)
    answers, warning = expected_answers(report, tmp_path)
    assert len(answers) == 60 and warning is None
    (tmp_path / "matching.json").write_bytes(dataset + b" ")
    answers, warning = expected_answers(report, tmp_path)
    assert answers == {} and "SHA-256" in warning


# 做什么：验证报告中的任意路径不会被直接打开。
# 为什么需要：报告展示不应变成读取其他文件的入口。
def test_report_dataset_path_is_not_used_to_read_arbitrary_files(report, tmp_path):
    report.dataset_path = str(Path(".env").resolve())
    assert expected_answers(report, tmp_path)[0] == {}


# 做什么：验证用于展示答案的数据集也拒绝重复题号。
# 为什么需要：避免字典覆盖导致标准答案对应错误。
def test_duplicate_dataset_ids_are_not_silently_accepted(report, tmp_path):
    payload = json.loads(Path("data/datasets/ecommerce_eval_v1.json").read_text(encoding="utf-8"))
    payload[1]["id"] = payload[0]["id"]
    content = json.dumps(payload).encode()
    (tmp_path / "bad.json").write_bytes(content)
    report.dataset_sha256 = hashlib.sha256(content).hexdigest()
    answers, warning = expected_answers(report, tmp_path)
    assert not answers and warning


# 做什么：验证状态、幻觉与类别筛选返回正确题目。
# 为什么需要：用户筛选问题时不能漏掉或混入不符结果。
def test_status_hallucination_and_category_filters(report):
    report.case_results[0].status = "TIMEOUT"
    report.case_results[0].judge_result = None
    report.case_results[1].judge_result.hallucination = True
    report.case_results[1].judge_result.unsupported_claims = ["fixture claim"]
    assert filter_cases(report, status="FAIL") == [report.case_results[0]]
    assert len(filter_cases(report, status="SUCCESS")) == 59
    assert filter_cases(report, hallucination="是") == [report.case_results[1]]
    assert filter_cases(report, hallucination="未评判") == [report.case_results[0]]
    assert all(r.category == "normal" for r in filter_cases(report, category="normal"))
    assert filter_cases(report, category="absent") == []


# 做什么：验证报告目录只收录合法完整报告。
# 为什么需要：快照和坏文件不能作为正常报告展示。
def test_catalog_ignores_snapshots_and_invalid_reports(report, tmp_path):
    (tmp_path / "candidate.json").write_bytes(json_bytes(report))
    (tmp_path / "baseline.json").write_text("{}", encoding="utf-8")
    assert list(report_catalog(tmp_path)) == [str(tmp_path / "candidate.json")]


# 做什么：验证导出前后原报告内容相同。
# 为什么需要：格式化展示不能修改评测事实。
def test_export_does_not_mutate_report(report, tmp_path):
    before = json_bytes(report)
    export_report(report, tmp_path)
    assert before == json_bytes(report)


# 做什么：验证零题报告仍能导出合法空表和模式说明。
# 为什么需要：空结果也需要可查看而非程序崩溃。
def test_empty_report_exports_header_only_tables(report):
    report.case_results = []
    report.summary = aggregate_results([])
    book = load_workbook(BytesIO(excel_bytes(report)))
    assert book["Case Results"].max_row == 4
    assert "JUDGE MODE: MOCK" in mode_notice(report)
