# 文件作用：把已有报告导出成完整 JSON 和四张 Excel 工作表。
# 为什么有它：方便分享与查看结果，同时处理长文本、非法字符和公式样式字符串。
"""Portable Python JSON/XLSX export of existing reports, with no metric calculation."""

from io import BytesIO
from pathlib import Path
import os
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.evaluation.models import EvaluationReport
from .data import as_json, expected_answers, mode_notice


# 做什么：将完整报告编码成 UTF-8 JSON 字节。
# 为什么需要：网页下载与文件导出可以复用同一内容。
def json_bytes(report: EvaluationReport) -> bytes:
    return report.model_dump_json(indent=2).encode("utf-8")


# 做什么：转义 Excel 不接受的控制字符，并标明超长文本截断。
# 为什么需要：避免导出文件损坏，完整原文仍保留在 JSON。
def _text(value: str) -> str:
    # Excel/XML cannot store these controls. JSON remains the lossless source.
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda m: f"\\u{ord(m[0]):04x}", value)
    return value if len(value) <= 32767 else value[:32720] + " [TRUNCATED: see JSON export]"


# 做什么：创建并格式化一张工作表，把文本明确写成字符串。
# 为什么需要：让结果可读，并防止以等号开头的文本被当成公式。
def _sheet(workbook, title, headers, rows, notice):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.append([title])
    sheet.append([notice])
    sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(value) if isinstance(value, str) else value for value in row])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = _text(cell.value)
                cell.data_type = "s"  # Literal strings, including '=...', never spreadsheet formulas.
            cell.font = Font(name="Arial", size=11, color="20302C")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                cell.number_format = "0.0000" if isinstance(cell.value, float) else "0"
                cell.alignment = Alignment(horizontal="right", vertical="top")
    sheet["A1"].alignment = Alignment(wrap_text=False)
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color="20302C")
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="755B17")
    # Note is repeated in an ordinary cell; detailed fields below remain filterable.
    sheet.row_dimensions[2].height = 65
    for cell in sheet[4]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285C4D")
    sheet.row_dimensions[4].height = 34
    for i, header in enumerate(headers, 1):
        width = 52 if header in {"query", "expected_answer", "rag_answer", "retrieved_chunks", "errors", "reason", "raw_output", "value"} else max(20, min(30, len(header) + 3))
        sheet.column_dimensions[get_column_letter(i)].width = width
    # A spans the note visually through empty cells; wrap only the note via a generous row height.
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, min(6, len(headers))))
    if title == "Summary":
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 76
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, sheet.max_row)}"
    for n in range(5, sheet.max_row + 1):
        sheet.row_dimensions[n].height = 60 if title != "Summary" else 30
    return sheet


# 做什么：把版本、汇总、逐题、检索和裁判结果写入四张工作表。
# 为什么需要：提供便于人工查看的 Excel，同时保留缺失值和模式提示。
def excel_bytes(report: EvaluationReport, answers: dict[str, str] | None = None,
                answer_warning: str | None = None) -> bytes:
    answers = answers or {}
    workbook = Workbook()
    workbook.remove(workbook.active)
    notice = mode_notice(report) + " 空白表示缺失；长文本可在编辑栏查看，超过32767字符时截断，完整内容见JSON。"
    if answer_warning:
        notice += " " + answer_warning
    summary = []
    for key in ("timestamp", "dataset_version", "dataset_sha256", "corpus_sha256", "embedding_model",
                "rag_model", "top_k", "judge_model", "judge_prompt_version", "rag_mode", "judge_mode",
                "quality_metrics_are_synthetic"):
        value = getattr(report, key)
        summary.append(["Metadata", key, value.isoformat() if key == "timestamp" else value])
    summary.append(["Metadata", "corpus_version", "sha256:" + report.corpus_sha256])
    for section, values in report.summary.model_dump().items():
        for key, value in values.items():
            summary.append([section, key, as_json(value) if isinstance(value, dict) else value])
    _sheet(workbook, "Summary", ["section", "metric", "value"], summary, notice)
    case_rows, retrieval_rows, judge_rows = [], [], []
    for row in report.case_results:
        rag, retrieval, judge = row.rag_result, row.retrieval_metrics, row.judge_result
        case_rows.append([
            row.case_id, row.category, row.answerable, row.status, row.query, answers.get(row.case_id),
            rag.rag_answer if rag else None,
            as_json([doc.model_dump() for doc in rag.retrieved_docs]) if rag else None,
            row.error_type, as_json([err.model_dump(mode="json") for err in row.errors]),
            row.latency_ms, row.retry_count, report.rag_mode, report.judge_mode,
        ])
        retrieval_rows.append([
            row.case_id, row.answerable, report.top_k, retrieval.recall_at_k if retrieval else None,
            retrieval.precision_at_k if retrieval else None, retrieval.rr if retrieval else None,
            as_json(retrieval.retrieved_chunk_ids) if retrieval else None,
            as_json(retrieval.relevant_chunk_ids) if retrieval else None,
            "EVALUATED" if retrieval and row.answerable else ("EXCLUDED" if not row.answerable else "UNAVAILABLE"),
            report.rag_mode,
        ])
        judge_rows.append([
            row.case_id, "AVAILABLE" if judge else "UNAVAILABLE", report.judge_mode,
            judge.answer_correctness if judge else None, judge.faithfulness if judge else None,
            judge.answer_relevance if judge else None, judge.completeness if judge else None,
            judge.hallucination if judge else None, as_json(judge.unsupported_claims) if judge else None,
            judge.reason if judge else None, judge.refusal_detected if judge else None,
            row.refusal_correct, judge.metadata.judge_model if judge else report.judge_model,
            judge.metadata.prompt_version if judge else report.judge_prompt_version,
            judge.metadata.latency_ms if judge else None, judge.metadata.raw_output if judge else None,
        ])
    _sheet(workbook, "Case Results", ["case_id", "category", "answerable", "status", "query", "expected_answer",
           "rag_answer", "retrieved_chunks", "error_type", "errors", "latency_ms", "retry_count", "rag_mode", "judge_mode"], case_rows, notice)
    _sheet(workbook, "Retrieval Metrics", ["case_id", "answerable", "top_k", "recall_at_k", "precision_at_k", "rr",
           "retrieved_chunk_ids", "relevant_chunk_ids", "availability", "rag_mode"], retrieval_rows, notice)
    _sheet(workbook, "Judge Results", ["case_id", "availability", "judge_mode", "correctness", "faithfulness", "relevance",
           "completeness", "hallucination", "unsupported_claims", "reason", "refusal_detected", "refusal_correct",
           "judge_model", "prompt_version", "latency_ms", "raw_output"], judge_rows, notice)
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


# 做什么：先写临时二进制文件，再替换目标文件。
# 为什么需要：降低导出时留下半成品的风险。
def _atomic_bytes(path: Path, content: bytes):
    path.parent.mkdir(parents=True, exist_ok=True)
    name = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".tmp", delete=False) as handle:
            name = handle.name
            handle.write(content)
        os.replace(name, path)
    finally:
        if name and Path(name).exists():
            Path(name).unlink()


# 做什么：匹配标准答案并保存 JSON 与 Excel 两种报告。
# 为什么需要：一次完成已有评测结果的分享与导出。
def export_report(report: EvaluationReport, output: str | Path,
                  dataset_dir: str | Path = "data/datasets") -> tuple[Path, Path]:
    output = Path(output)
    answers, warning = expected_answers(report, dataset_dir)
    # Prepare both formats before writing; each file replacement is atomic.
    json_data, xlsx_data = json_bytes(report), excel_bytes(report, answers, warning)
    json_path, xlsx_path = output / "evaluation_report.json", output / "evaluation_report.xlsx"
    _atomic_bytes(json_path, json_data)
    _atomic_bytes(xlsx_path, xlsx_data)
    return json_path, xlsx_path
